"""
ingest.py

Data ingestion module for a RAG (Retrieval-Augmented Generation) pipeline.

Supports pulling raw content from three source types and normalizing them
into LangChain `Document` objects:
    1. YouTube videos / playlists (via yt-dlp + youtube-transcript-api)
    2. Local PDF files (via pypdf)
    3. Local PowerPoint files (via python-pptx)

NOTE: This module is intentionally scoped to ingestion only. Chunking,
embedding, and vector-store persistence are handled downstream.
"""

from __future__ import annotations

import logging
import time
from pathlib import Path
from typing import List, Optional

from dotenv import load_dotenv
from langchain_core.documents import Document

# Third-party ingestion dependencies
import yt_dlp
from youtube_transcript_api import (
    YouTubeTranscriptApi,
    TranscriptsDisabled,
    NoTranscriptFound,
)
from pypdf import PdfReader
from pptx import Presentation

# Load environment variables from a .env file (API keys, config, etc.)
load_dotenv()

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
)
logger = logging.getLogger(__name__)


class DocumentIngestor:
    """
    Handles ingestion of raw source material (YouTube, PDF, PPTX) into a
    list of LangChain `Document` objects, ready for downstream chunking
    and embedding.
    """

    # ------------------------------------------------------------------ #
    # YouTube
    # ------------------------------------------------------------------ #
    def ingest_youtube(self, url: str) -> List[Document]:
        """
        Ingest a YouTube URL, which may point to a single video or a playlist.

        For each video found, fetches its transcript (if available) and
        wraps it in a `Document` with source URL and title metadata.

        Args:
            url: A YouTube video or playlist URL.

        Returns:
            A list of `Document` objects, one per video with a retrievable
            transcript. Videos without transcripts are skipped (and logged).
        """
        video_entries = self._resolve_youtube_entries(url)

        documents: List[Document] = []
        for entry in video_entries:
            video_id = entry.get("id")
            title = entry.get("title", "Unknown title")
            video_url = f"https://www.youtube.com/watch?v={video_id}"

            transcript_text = self._fetch_youtube_transcript(video_id)
            if transcript_text is None:
                logger.warning("Skipping video with no transcript: %s (%s)", title, video_url)
                continue

            documents.append(
                Document(
                    page_content=transcript_text,
                    metadata={
                        "source": video_url,
                        "title": title,
                        "video_id": video_id,
                        "type": "youtube",
                    },
                )
            )
            
            # Pause for 5 seconds between videos to avoid getting IP-blocked by YouTube
            time.sleep(5)

        logger.info("Ingested %d YouTube document(s) from %s", len(documents), url)
        return documents

    def _resolve_youtube_entries(self, url: str) -> List[dict]:
        """
        Use yt-dlp in metadata-only mode to determine whether `url` is a
        single video or a playlist, and return a flat list of video entries
        (each with at least 'id' and 'title').
        """
        ydl_opts = {
            "quiet": True,
            "no_warnings": True,
            "extract_flat": True,  # don't download, just pull metadata
            "skip_download": True,
        }

        with yt_dlp.YoutubeDL(ydl_opts) as ydl:
            info = ydl.extract_info(url, download=False)

        # Playlists expose their videos under "entries"; a single video does not.
        if info.get("entries"):
            entries = [e for e in info["entries"] if e]
        else:
            entries = [info]

        return entries

    def _fetch_youtube_transcript(self, video_id: str) -> Optional[str]:
        """
        Fetch and concatenate the transcript for a single YouTube video ID.

        Returns None if no transcript is available (disabled or not found).
        """
        try:
            yt_api = YouTubeTranscriptApi()
            transcript_list = yt_api.fetch(video_id)
        except (TranscriptsDisabled, NoTranscriptFound) as e:
            logger.warning("No transcript available for video %s: %s", video_id, e)
            return None
        except Exception as e:
            logger.error("Unexpected error fetching transcript for %s: %s", video_id, e)
            return None

        # Bulletproof loop to handle both objects (new versions) and dictionaries (old versions)
        pieces = []
        for chunk in transcript_list:
            # If the chunk is an object (FetchedTranscriptSnippet)
            if hasattr(chunk, "text"):
                if chunk.text:
                    pieces.append(chunk.text)
            # If the chunk is a dictionary
            elif isinstance(chunk, dict):
                if chunk.get("text"):
                    pieces.append(chunk.get("text"))

        return " ".join(pieces)

    # ------------------------------------------------------------------ #
    # PDF
    # ------------------------------------------------------------------ #
    def ingest_pdf(self, file_path: str) -> List[Document]:
        """
        Extract text from a local PDF file, one `Document` per page.

        Args:
            file_path: Path to the local PDF file.

        Returns:
            A list of `Document` objects with page-level granularity.
            Blank pages are skipped.
        """
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"PDF file not found: {file_path}")

        reader = PdfReader(str(path))
        documents: List[Document] = []

        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            text = text.strip()
            if not text:
                logger.debug("Skipping empty page %d in %s", page_number, path.name)
                continue

            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": str(path),
                        "title": path.stem,
                        "page": page_number,
                        "total_pages": len(reader.pages),
                        "type": "pdf",
                    },
                )
            )

        logger.info("Ingested %d page(s) from PDF %s", len(documents), path.name)
        return documents

    # ------------------------------------------------------------------ #
    # PowerPoint
    # ------------------------------------------------------------------ #
    def ingest_ppt(self, file_path: str) -> List[Document]:
        """
        Extract text from a local .pptx file, one `Document` per slide.

        All text-bearing shapes on a slide (titles, body text, text boxes)
        are concatenated into a single slide-level document.

        Args:
            file_path: Path to the local .pptx file.

        Returns:
            A list of `Document` objects with slide-level granularity.
            Blank slides are skipped.
        """
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"PPTX file not found: {file_path}")

        presentation = Presentation(str(path))
        documents: List[Document] = []
        total_slides = len(presentation.slides)

        for slide_number, slide in enumerate(presentation.slides, start=1):
            slide_text = self._extract_slide_text(slide)
            if not slide_text:
                logger.debug("Skipping empty slide %d in %s", slide_number, path.name)
                continue

            documents.append(
                Document(
                    page_content=slide_text,
                    metadata={
                        "source": str(path),
                        "title": path.stem,
                        "slide": slide_number,
                        "total_slides": total_slides,
                        "type": "pptx",
                    },
                )
            )

        logger.info("Ingested %d slide(s) from PPTX %s", len(documents), path.name)
        return documents

    def _extract_slide_text(self, slide) -> str:
        """Concatenate text from all text-bearing shapes on a single slide."""
        pieces = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                shape_text = "\n".join(
                    paragraph.text for paragraph in shape.text_frame.paragraphs if paragraph.text
                )
                if shape_text.strip():
                    pieces.append(shape_text.strip())

            # Tables aren't text frames but often carry meaningful content.
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        pieces.append(row_text)

        return "\n".join(pieces).strip()


if __name__ == "__main__":
    # Minimal smoke-test / usage example.
    ingestor = DocumentIngestor()

    # Example (uncomment and adjust paths/URLs to test locally):
    # docs = ingestor.ingest_youtube("https://www.youtube.com/watch?v=VIDEO_ID")
    # docs = ingestor.ingest_pdf("example.pdf")
    # docs = ingestor.ingest_ppt("example.pptx")
    # for d in docs:
    #     print(d.metadata, "->", d.page_content[:80])
    pass